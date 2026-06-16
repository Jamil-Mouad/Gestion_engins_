"""
Generate test Excel files for the Import feature.
Run: python3 generate-test-files.py

IMPORTANT: Update VEHICULE_CODES below to match real vehicle codes in your database.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

HERE = Path(__file__).parent

# ── helpers ─────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="2563EB")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 12)


# ── interventions ────────────────────────────────────────────────────────────
# Update vehicule_code values to match real codes in your DB
VEHICULE_CODES = ["CAM-401", "CAM-403", "CAM-405"]

wb_int = openpyxl.Workbook()
ws_int = wb_int.active
ws_int.title = "Interventions"

int_headers = [
    "vehicule_code", "type_intervention", "categorie",
    "partie", "n_bs", "date_debut", "cout_pieces", "observations",
]
int_rows = [
    [VEHICULE_CODES[0], "Réparation", "corrective",  "Moteur",  "BS-2026-001", "2026-01-15", 1500, "Test import"],
    [VEHICULE_CODES[1], "Vidange",    "preventive", "Carter",  "BS-2026-002", "2026-02-10",  800, "Vidange moteur"],
    [VEHICULE_CODES[2], "Diagnostic", "controle",  "Freins",  "BS-2026-003", "2026-03-05",  200, "Bruit anormal"],
]
write_sheet(ws_int, int_headers, int_rows)
out_int = HERE / "interventions-test.xlsx"
wb_int.save(out_int)
print(f"Created: {out_int}")

# ── demandes d'achat ─────────────────────────────────────────────────────────

wb_da = openpyxl.Workbook()
ws_da = wb_da.active
ws_da.title = "Demandes d'achat"

da_headers = ["piece_libre", "reference", "quantite", "motif", "urgence", "demandeur_nom"]
da_rows = [
    ["Filtre à huile",          "FH-001",  2, "Remplacement périodique", "Haute",   "Ahmed Benali"],
    ["Courroie de distribution","CD-002",  1, "Usure constatée",          "Haute",   "Youssef Idrissi"],
    ["Plaquettes de frein",     "PF-003",  4, "Contrôle révision",        "Moyenne", "Mourad Tazi"],
]
write_sheet(ws_da, da_headers, da_rows)
out_da = HERE / "demandes-achat-test.xlsx"
wb_da.save(out_da)
print(f"Created: {out_da}")
